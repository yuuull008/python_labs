import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from src.lib.io_helpers import ensure_parent_dir


def csv_to_xlsx(csv_path: str | Path, xlsx_path: str | Path) -> None:
    """Конвертирует CSV в XLSX.
    ValueError: неверное расширение файла; CSV пустой или в нём только заголовок.
    FileNotFoundError: CSV не найден.
    """
    csv_p = Path(csv_path)
    xlsx_p = Path(xlsx_path)
    
    if not csv_p.exists():
        raise FileNotFoundError(f"CSV не найден: {csv_p}")
    if csv_p.suffix.lower() != ".csv":
        raise ValueError("Ожидался файл с расширением .csv")
    if xlsx_p.suffix.lower() != ".xlsx":
        raise ValueError("Ожидался файл с расширением .xlsx")

    with csv_p.open(encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))

    if len(rows) < 2:
        raise ValueError("CSV пустой или содержит только заголовок")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for row in rows:
        ws.append(row)

    column_count = max(len(row) for row in rows)
    for col_index in range(column_count):
        longest = 0
        for row in rows:
            if col_index < len(row):
                longest = max(longest, len(row[col_index]))
        letter = get_column_letter(col_index + 1)
        ws.column_dimensions[letter].width = max(8, longest)

    ensure_parent_dir(xlsx_p)
    wb.save(xlsx_p)


if __name__ == "__main__":
    csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
    csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")
    csv_to_xlsx("data/samples/products.csv", "data/out/products.xlsx")
    print("Готово: результаты в data/out/")